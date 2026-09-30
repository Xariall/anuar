import json, glob, sys, os, unicodedata
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from build_matrix import DEF, objs
from apply_maps import rebuild
import verify_quotes as vq
NAMES={'займ':'займ','залог':'залог','гарантия':'гарантия','поручение':'поручение','дкп':'ДКП_лизинг','лизинг':'финансовый_лизинг'}
TITLES={'займ':'Договор займа (кредита, микрокредита, бюджетного кредита)','залог':'Договор залога','гарантия':'Договор гарантии (по содержанию — поручительство)','поручение':'Договор поручения (агентский)','дкп':'Договор купли-продажи для передачи в финансовый лизинг','лизинг':'Договор финансового лизинга'}
HDR=PatternFill('solid',fgColor='D9E1F2'); GREY=PatternFill('solid',fgColor='F2F2F2'); YEL=PatternFill('solid',fgColor='FFF2CC')
MAXC=32000
# Ручная сверка цитат, не найденных автосверкой (фрагменты найдены grep по _work/text/; расхождение только в оформлении)
MANUAL_OK=('Образец ДКП лизинг','Область Жетісу/договора/Договор займа','КД Балмұқан','договор займа (с переводом УСХ)','Область Абай/договора/Договор залога')
MANUAL_NOTE='подтверждена вручную: текст есть в источнике; отличие в маркерах списка, переносах строк или длине подчёркиваний'
def ic_note(desc):
    d=desc.lower()
    if 'солидарн' in d and 'назван' in d and 'залогодерж' not in d:
        return 'Несоответствие названия и содержания (гарантия / поручительство). По правилам заказчика это не противоречие; вопрос квалификации — спорный случай №1 этапа 1, решает юрист.'
    return ''
def cell_text(items):
    parts=[]
    for it in items:
        s=it['summary'].strip()
        if s=='отсутствует': parts.append('отсутствует'); continue
        parts.append(f"• {s} [{it['ref']}]")
    t='\n'.join(parts)
    return t if len(t)<=MAXC else t[:MAXC]+' …[обрезано: см. JSON]'
def load_analysis(kind):
    a={}
    for fn in sorted(glob.glob(f'analysis_{kind}_*.json')):
        a.update(json.load(open(fn)))
    return a
def fmt(ws,widths,freeze='B2'):
    for c in ws[1]: c.font=Font(bold=True); c.fill=HDR; c.alignment=Alignment(wrap_text=True,vertical='top')
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes=freeze
def sheet_list(wb,title,header,rows,widths):
    ws=wb.create_sheet(title); ws.append(header)
    for r in rows: ws.append([ (x if not isinstance(x,str) or len(x)<=MAXC else x[:MAXC]) for x in r])
    fmt(ws,widths,'A2'); return ws
def build(kind, extra_objs=(), appx=None):
    m=rebuild(kind); an=load_analysis(kind); cols=m['columns']
    wb=Workbook(); ws=wb.active; ws.title='Сравнение'
    ws.append([TITLES[kind]+' — сравнение условий по договорам. В ячейке: краткое содержание [пункт источника]. «отсутствует» — условие в договоре не найдено.'])
    ws.append(['Условие / раздел']+[f"{c['label']}\n{c['file'].split('/')[-1]}" for c in cols]+['Различия','Противоречия'])
    for t,cells in m['rows'].items():
        a=an.get(t,{})
        ws.append([t]+[cell_text(cells[c['label']]) if c['label'] in cells else '—' for c in cols]+[a.get('различия','(не заполнено)'),a.get('противоречия','(не заполнено)')])
    for c in ws[2]: c.font=Font(bold=True); c.fill=HDR; c.alignment=Alignment(wrap_text=True,vertical='top')
    for row in ws.iter_rows(min_row=3):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
        row[0].font=Font(bold=True)
    ws.column_dimensions['A'].width=34
    for i in range(len(cols)): ws.column_dimensions[get_column_letter(i+2)].width=42
    n=len(cols); ws.column_dimensions[get_column_letter(n+2)].width=70; ws.column_dimensions[get_column_letter(n+3)].width=70
    ws.freeze_panes='B3'
    # source objects for side sheets
    seen=[]; 
    for c in DEF[kind]:
        o=c['obj']
        if o not in seen: seen.append(o)
    for o in extra_objs:
        if o not in seen: seen.append(o)
    ic=[];kd=[];rs=[];st=[]
    for o in seen:
        f=o['file']; src=o['_src']
        for x in o.get('internal_contradictions',[]):
            desc=x.get('description',str(x)) if isinstance(x,dict) else str(x)
            ic.append([f,', '.join(map(str,x.get('refs',[]))) if isinstance(x,dict) else '',desc,ic_note(desc)])
        for x in o.get('kk_divergences',[]):
            kd.append([f,x.get('ref_rus',''),x.get('ref_kk',''),x.get('description','')] if isinstance(x,dict) else [f,'','',str(x)])
        for x in o.get('residuals',[]):
            rs.append([f,x.get('ref','') if isinstance(x,dict) else '',(x.get('description') or x.get('text') or json.dumps(x,ensure_ascii=False)) if isinstance(x,dict) else str(x)])
        st.append([f,str(o.get('status','')),str(o.get('notes',''))])
    sheet_list(wb,'Внутренние противоречия',['Договор','Пункты','Описание','Примечание (квалификация)'],ic,[60,45,110,50])
    sheet_list(wb,'Расхождения каз.-рус.',['Договор','Пункт (рус.)','Пункт (каз.)','Описание'],kd,[60,30,30,110])
    sheet_list(wb,'Остатки чужих данных',['Договор','Пункт','Описание'],rs,[60,40,110])
    sheet_list(wb,'Статус и полнота чтения',['Договор','Статус','Примечания экстрактора (что прочитано)'],st,[60,40,120])
    # quote check
    qc=[]
    for o in seen:
        c=vq.find_txt(o['file'])
        if not c: continue
        own=vq.texts[c[0]]
        for it in o['items']:
            q=it.get('quote','')
            if q and it.get('summary')!='отсутствует' and vq.norm(q) not in own:
                short=('...' in q or '…' in q)
                manual=MANUAL_NOTE if (not short and any(m in o['file'] for m in MANUAL_OK)) else ''
                qc.append([o['file'],it.get('ref',''),q[:300],'сокращённая цитата (…)' if short else 'не найдена дословно',manual])
    sheet_list(wb,'Проверка цитат',['Договор','Пункт','Цитата (выдержка агента)','Результат автосверки','Ручная сверка'],qc,[60,45,90,30,50])
    if appx: appx(wb)
    out=f'../02_сравнение_{NAMES[kind]}.xlsx'; wb.save(out); print(out,len(m['rows']),'rows',len(cols),'cols; ic',len(ic),'kk',len(kd),'res',len(rs),'quotes-flag',len(qc))
if __name__=='__main__':
    kinds=sys.argv[1:] or list(NAMES)
    for k in kinds:
        extra=()
        appx=None
        if k=='поручение':
            extra=[o for o in objs if o['_src']=='s2_A1.json']
            def appx(wb):
                rows=[]
                for o in objs:
                    if o['_src'] in ('s2_apa.json','s2_apb.json'):
                        for it in o['items']: rows.append([o['file'].split('/')[-1],'Кызылорда (отдельный файл)',it.get('part',''),it.get('topic',''),it.get('summary',''),it.get('ref','')])
                    if o['_src']=='s2_n43.json':
                        for it in o['items']:
                            if it.get('part','')!='Договор поручения': rows.append([o['file'].split('/')[-1],'Мангистау (внутри договора)',it.get('part',''),it.get('topic',''),it.get('summary',''),it.get('ref','')])
                    if o['_src']=='s2_A1.json':
                        for it in o['items']:
                            if it.get('part','').startswith('Приложение'): rows.append([o['file'].split('/')[-1],'Кызылорда (в договоре)',it.get('part',''),it.get('topic',''),it.get('summary',''),it.get('ref','')])
                ws=sheet_list(wb,'Приложения (СОМНЕВАЕМСЯ)',['Файл','Где находится','Приложение','Тема','Содержание','Ссылка'],rows,[40,28,30,30,100,50])
                ws.insert_rows(1); ws['A1']='ЯРЛЫК ЗАКАЗЧИКА: «Сомневаемся» — неясно, считать ли приложения к договору поручения отдельной группой или частью договора. Разобраны как приложения; решение не принято.'
                ws['A1'].fill=YEL
        build(k,extra,appx)
