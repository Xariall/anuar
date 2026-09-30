import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from collections import OrderedDict
rows=[]
for k in "ABCD": rows+=json.load(open(f"tmp/stage1_{k}.json"))
T={ 'Договор займа (кредита, микрокредита, бюджетного кредита)':[1,2,3,10,11,14,18,21,24,29,36,46,47,48,52,54,55,57],
 'Договор залога':[0,6,7,8,9,15,16,17,20,26,28,32,33,37,44,49,51,53,58,59],
 'Договор гарантии (по содержанию — поручительство)':[4,5,12,13,19,25,27,45,50,56],
 'Договор поручения (агентский): МИО — АО-агент':[34,35,42,43],
 'Договор купли-продажи для передачи в финансовый лизинг':[22,30,62,63],
 'Договор финансового лизинга':[31,64,65],
 'Приложение к договору поручения (не самостоятельный договор) — СОМНЕВАЕМСЯ':[38,39,40,41],
 'Составной файл (несколько договоров в одном файле)':[23,60,61]}
typ={i:t for t,l in T.items() for i in l}; assert len(typ)==66
def status(i,r):
    s=r['status'].lower()
    if i in (7,8): return 'версия'
    if s.startswith('дубликат'): return 'дубликат'
    if s.startswith('заполненный'): return 'заполненный экземпляр'
    return 'шаблон'
def j(x): 
    if isinstance(x,list): return '; '.join(str(a if not isinstance(a,dict) else a.get('name','')+' ['+str(a.get('location',''))+']') for a in x)
    return str(x)
wb=Workbook(); ws=wb.active; ws.title='Реестр'
H=['№','регион','путь к файлу','название','язык','статус','фактический тип (сводный)','фактический тип (по содержанию, детально)','договоры внутри файла','назначение','связанные договоры','кандидаты на объединение / признаки, препятствующие объединению','подтверждение статуса','заемщик/сторона','вид имущества (залог)','спорная классификация','причина спора']
ws.append(H)
for i,r in enumerate(rows):
    ws.append([i+1,r['region'],r['path'],r['title'],r['language'],status(i,r),typ[i],r['actual_type'],j(r['contracts_inside']),r['purpose'],j(r['related']),j(r['merge_candidates']),r['status_evidence']+' | '+'; '.join(map(str,r['evidence'])) if isinstance(r['evidence'],list) else r['status_evidence'],r['borrower_category'],r['property_kind'],'да' if r['disputed'] else 'нет',r['disputed_reason'] or ''])
hdr=PatternFill('solid',fgColor='DDDDDD')
def fmt(ws,widths):
    for c in ws[1]: c.font=Font(bold=True); c.fill=hdr; c.alignment=Alignment(wrap_text=True,vertical='top')
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
    for k,w in widths.items(): ws.column_dimensions[k].width=w
    ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
fmt(ws,{'A':5,'B':22,'C':45,'D':30,'E':16,'F':16,'G':30,'H':60,'I':50,'J':40,'K':45,'L':60,'M':60,'N':20,'O':20,'P':10,'Q':45})
s=wb.create_sheet('Сводка по типам'); s.append(['фактический тип','всего файлов','уникальных (без дубликатов)','дубликатов','заполненных экземпляров','регионы','файлы (№ в реестре)'])
for t,l in T.items():
    st=[status(i,rows[i]) for i in l]
    reg=sorted({rows[i]['region'].split('.')[0] for i in l},key=int)
    s.append([t,len(l),len(l)-st.count('дубликат'),st.count('дубликат'),st.count('заполненный экземпляр'),', '.join(reg),', '.join(str(i+1) for i in l)])
s.append(['ИТОГО',66,66-sum(1 for i in range(66) if status(i,rows[i])=='дубликат'),sum(1 for i in range(66) if status(i,rows[i])=='дубликат'),sum(1 for i in range(66) if status(i,rows[i])=='заполненный экземпляр'),'',''])
fmt(s,{'A':55,'B':12,'C':16,'D':12,'E':16,'F':40,'G':60}); s.freeze_panes='A2'
d=wb.create_sheet('Спорные случаи'); d.append(['№','путь','статус','тип (сводный)','причина'])
for i,r in enumerate(rows):
    if r['disputed']: d.append([i+1,r['path'],status(i,r),typ[i],r['disputed_reason']])
fmt(d,{'A':5,'B':55,'C':16,'D':40,'E':100}); d.freeze_panes='A2'
u=wb.create_sheet('Дубликаты и версии'); u.append(['№','путь','статус','подтверждение'])
for i,r in enumerate(rows):
    if status(i,r) in('дубликат','версия'): u.append([i+1,r['path'],status(i,r),r['status_evidence']])
fmt(u,{'A':5,'B':55,'C':12,'D':120}); u.freeze_panes='A2'
m=wb.create_sheet('Методика и ограничения')
for line in ['Источник: 66 файлов в шаблоны МИО/ (65 .doc/.docx + 1 .xlsx). Исходники не изменялись.',
 'Текстовые копии: LibreOffice в системе отсутствует, конвертация выполнена macOS textutil (.doc/.docx → .txt, UTF-8) в _work/text/ (65 файлов). .xlsx прочитан отдельно (8 листов).',
 'Классификация по предмету и ролям сторон, а не по названию файла; сводный тип присвоен при сборке реестра.',
 'Побайтовые дубликаты (md5): Туркестан — 3 файла в корне = 3 файла в договара/; Мангистау — «рус (1)» = «рус»; СКО — залог «(1)» = залог. По содержанию (текст идентичен, md5 различается): Павлодар — «Шаблон Договор займа.doc» = «… Павлодар.docx».',
 'Версии (тексты различаются): Актобе — «Договор залога.docx» и «Договор залога недвижимого имущества.docx» (~107 строк diff).',
 'ГЛУБИНА ЧТЕНИЯ: файлы читали 4 субагента. Полностью прочитаны не все: регион 1 (залог, займ), частично 2 (казахский блок кредита ИП), 5–6 (кроме ДК ИП и ДЗ ИП — частично), 8 «гарант», регионы 10–13 (заявлено «целиком»). Остальные — заголовки, пункты о сторонах и предмете, заключительные разделы, сравнение соседних файлов. Для типизации этого достаточно; для этапа 2 (сравнение условий) потребуется построчное чтение.',
 'Идентичность казахской и русской половин двуязычных файлов на этапе 1 не сверялась построчно (только на этапе 2 по правилам).',
 'Утверждения в столбцах «подтверждение» и «детально» — выдержки субагентов с указанием пунктов; выборочная проверка ссылок запланирована на этапе 5.']:
    m.append([line])
m.column_dimensions['A'].width=150
for row in m.iter_rows():
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
wb.save('01_реестр.xlsx')
for r in s.iter_rows(values_only=True): print(r)
print([ (i+1,rows[i]['path'].split('/')[-1][:30]) for i in range(66) if rows[i]['disputed']])
